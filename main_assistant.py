# main_assistant.py
import re
import json
import logging
import tkinter as tk
from tkinter import ttk
import asyncio
import docx
import PyPDF2
import mimetypes
import google.generativeai as genai
from google.generativeai.types import HarmCategory, HarmBlockThreshold
from main_assistant_prompts import agent_prompts, system_prompt
from command_execution import *
from memory_main import MemoryManager
from json_fixer_prompt import json_fixer_prompt
from length_limiter import LengthLimiter

# Configuration
config_path = os.path.join(os.path.dirname(__file__), 'config', 'config.json')
genai.configure(api_key=json.loads(open(config_path).read())['api_keys']['GEMINI_API_KEY'])

initialized_agents = {}

async def show_confirmation_dialog(agent, function_name, parameters):
    root = tk.Tk()
    root.title("Confirm Action")
    root.geometry("400x300") 
    
    try:
        frame = ttk.Frame(root, padding="10")
        frame.pack(fill=tk.BOTH, expand=True)

        ttk.Label(frame, text=f"The {agent} wants to execute:", wraplength=380).pack()
        ttk.Label(frame, text=f"Function: {function_name}", wraplength=380).pack()

        text_frame = ttk.Frame(frame)
        text_frame.pack(fill=tk.BOTH, expand=True, pady=5)

        text_area = tk.Text(text_frame, wrap=tk.WORD, width=40, height=8)
        text_area.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        
        scrollbar = ttk.Scrollbar(text_frame, orient=tk.VERTICAL, command=text_area.yview)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        
        text_area.config(yscrollcommand=scrollbar.set)
        
        params_str = json.dumps(parameters, indent=2)
        text_area.insert(tk.END, f"Parameters:\n{params_str}")
        text_area.config(state=tk.DISABLED)

        ttk.Label(frame, text="Do you allow this action?").pack(pady=5)

        result = asyncio.Future()

        def on_allow():
            result.set_result(True)
            root.quit()

        def on_pass():
            result.set_result(False)
            root.quit()

        button_frame = ttk.Frame(frame)
        button_frame.pack(pady=10)

        ttk.Button(button_frame, text="Allow", command=on_allow).pack(side=tk.LEFT, padx=10)
        ttk.Button(button_frame, text="Pass", command=on_pass).pack(side=tk.LEFT, padx=10)

        root.protocol("WM_DELETE_WINDOW", on_pass)
        root.mainloop()
        
        return await result
    except Exception as e:
        print(f"Error in show_confirmation_dialog: {str(e)}")
    finally:
        root.destroy()


async def execute_agent_instruction(agent_response, agent_name, config):
    try:
        parsed_response = json.loads(agent_response)
        
        function_name = parsed_response.get('function')
        parameters = parsed_response.get('parameters', {})
        
        current_mode = config["gui_settings"]["mode_selected"]
        
        if current_mode == "safe":
            user_confirmed = await show_confirmation_dialog(agent_name, function_name, parameters)
            if not user_confirmed:
                return "Action cancelled by user."
        else:
            user_confirmed = True
        
        agent_map = {   
            'CommandExecutionAgent': CommandExecutionAgent,
            'FileOperationsAgent': FileOperationsAgent,
            'WebInteractionAgent': WebInteractionAgent,
            'SystemManagementAgent': SystemManagementAgent,
            'GUIAutomationAgent': GUIAutomationAgent,
            'ProcessManagementAgent': ProcessManagementAgent,
            'NetworkOperationsAgent': NetworkOperationsAgent,
            'WindowManagementAgent': WindowManagementAgent,
            'MediaProcessingAgent': MediaProcessingAgent,
            'SpeechAgent': SpeechAgent,
            'DatabaseOperationsAgent': DatabaseOperationsAgent,
            'ClipboardAgent': ClipboardAgent,
            'RegistryAgent': RegistryAgent,
            'SecurityOperationsAgent': SecurityOperationsAgent,
            'VirtualMachineAgent': VirtualMachineAgent,
            'FinanceAgent': FinanceAgent,
            'WeatherAgent': WeatherAgent,
            'AudioOperationsAgent': AudioOperationsAgent,
            'TeacherAgent': TeacherAgent
        }
        
        agent_class = agent_map.get(agent_name)
        
        if agent_class is None:
            raise ValueError(f"No agent found for: {agent_name}")
        
        if agent_name not in initialized_agents:
            agent_instance = agent_class()
            
            if hasattr(agent_instance, 'initialize') and callable(getattr(agent_instance, 'initialize')):
                init_method = getattr(agent_instance, 'initialize')
                if asyncio.iscoroutinefunction(init_method):
                    await init_method()
                else:
                    init_method()
            
            initialized_agents[agent_name] = agent_instance
        else:
            agent_instance = initialized_agents[agent_name]
        
        func = getattr(agent_instance, function_name)
    
        if asyncio.iscoroutinefunction(func):
            result = await func(**parameters)
        else:
            result = func(**parameters)
        
        print(f"Agent {agent_name} executed function {function_name} with parameters: {parameters}")
        print(f"Result: {result}")
        
        return result
    
    except json.JSONDecodeError:
        return "Error: Invalid JSON response from agent"
    except AttributeError:
        return f"Error: Function {function_name} not found in agent {agent_name}"
    except Exception as e:
        return f"Error executing agent instruction: {str(e)}"

async def cleanup_agents():
    for agent_name, agent_instance in initialized_agents.items():
        if hasattr(agent_instance, 'close') and callable(getattr(agent_instance, 'close')):
            close_method = getattr(agent_instance, 'close')
            if asyncio.iscoroutinefunction(close_method):
                await close_method()
            else:
                close_method()
    initialized_agents.clear()

class MainAssistant:
    def __init__(self, config):
        self.chat_history = []
        self.agent_prompts = agent_prompts
        self.task_context = {}
        self.model_name = "gemini-1.5-pro"
        self.uploaded_files = []
        self.background_tasks = {}
        self.configure_api()
        self.memory_manager = MemoryManager()
        self.json_fixer_model = genai.GenerativeModel('gemini-1.5-flash', system_instruction=json_fixer_prompt)
        self.json_fixer_chat = self.json_fixer_model.start_chat(history=[])
        self.config = config
        self.length_limiter = LengthLimiter(os.environ.get("GEMINI_API_KEY"))
        self.task_context_limit = 100000
        self.function_output_limit = 10000 

    def configure_api(self):
        config_path = os.path.join(os.path.dirname(__file__), 'config', 'config.json')
        genai.configure(api_key=json.loads(open(config_path).read())['api_keys']['GEMINI_API_KEY'])

    def set_model(self, model_name):
        self.model_name = model_name
    
    async def get_user_input(self, prompt):
        return input(prompt)
        
    def cleanup_uploaded_files(self):
        for file_name in self.uploaded_files:
            try:
                genai.delete_file(file_name)
                print(f"Deleted file: {file_name}")
            except Exception as e:
                print(f"Error deleting file {file_name}: {str(e)}")
        self.uploaded_files = []

    def reset_chat_history(self):
        chat_history_str = self.get_chat_history_as_string()
        self.chat_history = []
        self.task_context = {}
        self.cleanup_uploaded_files()
        self.update_external_knowledge()
        logging.info("Chat history reset and uploaded files deleted in MainAssistant")
        return chat_history_str
    
    def get_chat_history(self):
        return self.chat_history

    def parse_json_response(self, response):
        def fix_json_issues(bad_json):
            # Adding quotes around unquoted keys
            bad_json = re.sub(r'(?<!")(\b\w+\b)(?=\s*:)', r'"\1"', bad_json)
            # Escaping backslashes correctly
            bad_json = re.sub(r'\\(?![\\"])', r'\\\\', bad_json)
            # Removing invalid escape sequences
            bad_json = re.sub(r'\\[^\\"]', r'', bad_json)
            # Fixing missing commas between key-value pairs
            bad_json = re.sub(r'(?<=\w)(\s*)(?=")', r',\1', bad_json)
            # Fixing missing braces
            if bad_json.count('{') > bad_json.count('}'):
                bad_json += '}'
            if bad_json.count('[') > bad_json.count(']'):
                bad_json += ']'
            # Fixing trailing commas before closing braces
            bad_json = re.sub(r',(\s*[}\]])', r'\1', bad_json)
            # Ensuring colons are followed by a space
            bad_json = re.sub(r':(\S)', r': \1', bad_json)
            # Fixing common typos in JSON keys/values
            bad_json = re.sub(r'\bNone\b', 'null', bad_json)
            bad_json = re.sub(r'\bTrue\b', 'true', bad_json)
            bad_json = re.sub(r'\bFalse\b', 'false', bad_json)
            return bad_json

        cleaned_response = re.sub(r'^```json\n|```$', '', response).strip()
        cleaned_response = cleaned_response.replace("```", "")

        try:
            return json.loads(cleaned_response)
        except json.JSONDecodeError as e:
            print(f"Error parsing JSON: {str(e)}")
            print("Attempting to fix JSON...")
            cleaned_response1 = fix_json_issues(cleaned_response)
            try:
                return json.loads(cleaned_response1)
            except json.JSONDecodeError as e:
                print("Failed to fix JSON. Sending to JSON Fixer model...")
                return self.fix_json_with_model(cleaned_response, str(e))

    def fix_json_with_model(self, bad_json, error_message):
        max_attempts = 5
        attempt = 0

        while attempt < max_attempts:
            try:
                model_response = self.json_fixer_chat.send_message(
                    f"Fix this JSON:\n\n{bad_json}\n\nError: {error_message}",
                    generation_config=self.json_fixer_generation_config
                )
                fixed_json = model_response.text
                cleaned_fixed_json = re.sub(r'^```json\n|```$', '', fixed_json).strip()
                cleaned_fixed_json = cleaned_fixed_json.replace("```", "")
                
                parsed_json = json.loads(cleaned_fixed_json)
                return parsed_json
            except json.JSONDecodeError as e:
                attempt += 1
                error_message = str(e)
                if attempt == max_attempts:
                    raise ValueError(f"Failed to fix JSON after {max_attempts} attempts")

    @property
    def json_fixer_generation_config(self):
        return genai.types.GenerationConfig(
            temperature=1,
            top_p=1,
            top_k=1,
            max_output_tokens=8000,
        )

    async def process_file(self, file_path):
        mime_type, _ = mimetypes.guess_type(file_path)
        file_type = mime_type.split('/')[0] if mime_type else 'unknown'
        file_name = os.path.basename(file_path)

        if file_type in ['image', 'video', 'audio'] or mime_type in [
            'text/plain', 'text/html', 'text/css', 'text/javascript', 'application/x-javascript',
            'text/x-typescript', 'application/x-typescript', 'text/csv', 'text/markdown',
            'text/x-python', 'application/x-python-code', 'application/json', 'text/xml',
            'application/rtf', 'text/rtf'
        ]:
            file = await self.upload_and_process_file(file_path, file_name)
            return file
        elif file_path.endswith('.pdf'):
            return self.extract_text_from_pdf(file_path)
        elif file_path.endswith('.docx'):
            return self.extract_text_from_docx(file_path)
        else:
            return f"Unsupported file type: {file_name}"

    async def upload_and_process_file(self, file_path, file_name):
        file = genai.upload_file(path=file_path, display_name=file_name)
        print(f"Uploaded file: {file.name}")
        self.uploaded_files.append(file.name)
        
        while file.state.name == "PROCESSING":
            await asyncio.sleep(1)
            file = genai.get_file(file.name)

        if file.state.name == "FAILED":
            raise ValueError(file.state.name)
        
        print("File processing completed.")
        
        return file

    async def extract_text_from_file(self, file_path):
        mime_type, _ = mimetypes.guess_type(file_path)
        
        if mime_type == 'application/pdf':
            return self.extract_text_from_pdf(file_path)
        elif mime_type == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document':
            return self.extract_text_from_docx(file_path)
        else:
            with open(file_path, 'r', encoding='utf-8') as file:
                return file.read()

    def extract_text_from_pdf(self, file_path):
        with open(file_path, 'rb') as file:
            pdf_reader = PyPDF2.PdfReader(file)
            return '\n'.join([page.extract_text() for page in pdf_reader.pages])

    def extract_text_from_docx(self, file_path):
        doc = docx.Document(file_path)
        return '\n'.join([paragraph.text for paragraph in doc.paragraphs])
    
    def get_formatted_memory(self):
        memory_data = self.memory_manager.get_all_memory()
        formatted_memory = "YOUR MEMORY:\n\n"
        
        memory_types = [
            ("LONG TERM MEMORY:", "long_term_memory"),
            ("SHORT TERM MEMORY:", "short_term_memory"),
            ("ABILITIES MEMORY:", "abilities"),
            ("USER PREFERENCES MEMORY:", "user_preferences")
        ]
        
        for title, key in memory_types:
            formatted_memory += f"{title}\n"
            if memory_data[key]: 
                for item in memory_data[key]:
                    formatted_memory += json.dumps(item, indent=2)
                    formatted_memory += "\n" + "-"*80 + "\n"
            else:
                formatted_memory += "No memories stored.\n"
            formatted_memory += "\n"
        
        return formatted_memory


    async def generate_response(self, content, agent_output=None):
        prepared_content = self.prepare_content(content, agent_output)
        
        try:
            if self.model_name == "gemini-1.0-pro":
                return await self.generate_response_1_0(prepared_content)
            else:
                return await self.generate_response_1_5(prepared_content)
        except Exception as e:
            print(f"Error generating response: {str(e)}")
            return "I apologize, but I encountered an error while generating a response. Could you please try again?"

    def prepare_content(self, content, agent_output=None):
        prepared_content = content.copy()
        
        formatted_memory = self.get_formatted_memory()
        prepared_content.append(formatted_memory)
        
        if self.task_context:
            latest_output = next((v for k, v in self.task_context.items() if isinstance(v, dict) and v.get('is_latest')), None)
            task_context_str = json.dumps(self.task_context, indent=2)
            if latest_output:
                prepared_content.append(f"Task context (with latest output highlighted):\n{task_context_str}")
            else:
                prepared_content.append(f"Task context:\n{task_context_str}")
        
        print(f"Prepared content: {prepared_content}")
        return prepared_content
                
    async def generate_response_1_5(self, content):
        model = genai.GenerativeModel(
            model_name=self.model_name,
            generation_config=self.generation_config_1_5,
            system_instruction=system_prompt
        )
        chat = model.start_chat(history=self.chat_history)
        
        response = await chat.send_message_async(
            content=content,
            generation_config=self.generation_config_1_5,
            safety_settings=self.safety_settings
        )
        
        self.chat_history = chat.history
        print(f"Raw Response: {response.text}")
        return response.text

    async def generate_response_1_0(self, content):
        model = genai.GenerativeModel(
            model_name=self.model_name,
            generation_config=self.generation_config_1_0
        )
        chat = model.start_chat(history=self.chat_history)

        prompt = f"System: {system_prompt}\n User: {content[0]}"
        for item in content[1:]:
            if isinstance(item, genai.types.file_types.File):
                prompt += f"\n[Attached file: {item.display_name}]"
            else:
                prompt += f"\n{item}"
        
        response = await chat.send_message_async(
            content=prompt,
            generation_config=self.generation_config_1_0,
            safety_settings=self.safety_settings
        )
        
        self.chat_history = chat.history
        print(f"Raw Response: {response.text}")
        return response.text
    
    def get_external_knowledge(self):
        config_file = os.path.join(os.path.dirname(__file__), "config", "external_knowledge", "knowledge.json")
        if os.path.exists(config_file):
            with open(config_file, 'r') as f:
                config = json.load(f)
                if config.get("enabled", False):
                    knowledge = []
                    for knowledge_type in ["train", "observe", "external"]:
                        if config[knowledge_type]["enabled"]:
                            for file in config[knowledge_type]["files"]:
                                file_path = os.path.join(os.path.dirname(__file__), "config", "external_knowledge", file)
                                with open(file_path, 'r') as kf:
                                    knowledge.append(f"{knowledge_type.capitalize()} Knowledge - {file}:\n{kf.read()}\n")
                    return "\n".join(knowledge)
        return ""
    
    def update_external_knowledge(self):
        config_file = os.path.join(os.path.dirname(__file__), "config", "knowledge.json")
        knowledge_dir = os.path.join(os.path.dirname(__file__), "config", "external_knowledge")
        if os.path.exists(config_file):
            with open(config_file, 'r') as f:
                config = json.load(f)
                self.external_knowledge_enabled = config.get("enabled", False)
                self.external_knowledge = []
                if self.external_knowledge_enabled:
                    # Add train data if enabled
                    if config["train"]["enabled"]:
                        train_file = r'action_data\ability_memory.json'
                        if os.path.exists(train_file):
                            with open(train_file, 'r', encoding='utf-8') as tf:
                                self.external_knowledge.append({
                                    "type": "train",
                                    "file": "ability_memory.json",
                                    "content": json.load(tf)
                                })
                    
                    # Add observe data if enabled
                    if config["observe"]["enabled"]:
                        observe_file = r'observe_data\observed_summary.txt'
                        if os.path.exists(observe_file):
                            with open(observe_file, 'r', encoding='utf-8') as of:
                                self.external_knowledge.append({
                                    "type": "observe",
                                    "file": "observed_summary.txt",
                                    "content": of.read()
                                })
                    
                    # Add other external files
                    for file in os.listdir(knowledge_dir):
                        file_path = os.path.join(knowledge_dir, file)
                        if file.lower().endswith(('.txt', '.pdf', '.docx')):
                            try:
                                if file.lower().endswith('.txt'):
                                    with open(file_path, 'r', encoding='utf-8') as f:
                                        content = f.read()
                                elif file.lower().endswith('.pdf'):
                                    content = self.extract_text_from_pdf(file_path)
                                elif file.lower().endswith('.docx'):
                                    content = self.extract_text_from_docx(file_path)
                                
                                self.external_knowledge.append({
                                    "type": "external",
                                    "file": file,
                                    "content": content
                                })
                            except Exception as e:
                                print(f"Error processing file {file}: {str(e)}")
                else:
                    print("External Knowledge is disabled.")
        else:
            self.external_knowledge_enabled = False
            self.external_knowledge = []

    async def process_user_request(self, user_message, uploaded_files=None):
        try:
            print(f"Received uploaded files: {uploaded_files}")
            content = [f"User: {user_message}"]

            if uploaded_files:
                for file_name, file_content in uploaded_files:
                    if isinstance(file_content, genai.types.file_types.File):
                        content.append(file_content)
                    else:
                        # For large text files, truncate the content
                        max_chars = 50000
                        truncated_content = file_content[:max_chars]
                        if len(file_content) > max_chars:
                            truncated_content += "\n...[Content truncated due to length]..."
                        content.append(f"\n\n{'=' * 80}\nAttached file content: {file_name}\n\n{truncated_content}")
            
            # Add external knowledge
            if self.external_knowledge_enabled:
                for knowledge in self.external_knowledge:
                    content.append(f"\n\n{'=' * 80}\nExternal Knowledge ({knowledge['type']}) - {knowledge['file']}:\n\n")
                    if isinstance(knowledge['content'], str):
                        content.append(knowledge['content'])
                    else:
                        content.append(json.dumps(knowledge['content'], indent=2))

            if not self.task_context:
                current_subtask_id = 1
                print("Initializing current_subtask_id to 1 because task_context is empty.")
            else:
                current_subtask_id = max(int(k) for k in self.task_context.keys()) + 1
                print(f"Continuing with current_subtask_id: {current_subtask_id}")
            
            continue_execution = True

            initial_response = True
            while continue_execution:
                if initial_response:
                    content.append(f"\nCurrent Subtask ID(THE PROVIDED SUBTASK WITH THIS ID WILL BE EXECUTED): {current_subtask_id}")
                    main_response = await self.generate_response(content, None)
                    initial_response = False
                else:
                    # Check and limit task context length
                    task_context_str = json.dumps(self.task_context)
                    if len(task_context_str) > self.task_context_limit:
                        summarized_context = await self.length_limiter.summarize(task_context_str, self.task_context_limit, preserve_keys=['current_subtask_id'])
                        self.task_context = json.loads(summarized_context)

                    continue_task = [
                        "The previous step has been completed. Please analyze the result and determine the next action.",
                        f"Current Subtask ID(THE PROVIDED SUBTASK WITH THIS ID WILL BE EXECUTED): {current_subtask_id}",
                        f"Task context: {json.dumps(self.task_context)}"
                    ]
                    main_response = await self.generate_response(continue_task)

                parsed_response = self.parse_json_response(main_response)
                print(f"AI Response Type: {'Text Processing' if 'direct_response' in parsed_response else 'Task Breakdown'}")

                # Update plan information
                task_breakdown = parsed_response.get('task_breakdown', [])
                if task_breakdown:
                    current_task = next((task for task in task_breakdown if int(task['subtask_id']) == current_subtask_id), None)
                    if current_task:
                        plan_info = {
                            "current_plan": task_breakdown[0].get('subtask', 'N/A'),
                            "current_subtask": f"Subtask {current_subtask_id}: {current_task.get('instruction_to_agent', 'N/A')}"
                        }
                    else:
                        plan_info = {
                            "current_plan": "No specific plan needed",
                            "current_subtask": "Responding directly"
                        }
                else:
                    plan_info = {
                        "current_plan": "No specific plan needed",
                        "current_subtask": "Responding directly"
                    }
                self.update_plan_callback(plan_info)


                if not parsed_response or 'task_breakdown' not in parsed_response or not parsed_response['task_breakdown']:
                    return parsed_response.get('response_to_user', "I'm sorry, but I couldn't generate a proper response. Could you please try again?")

                current_task = next((task for task in parsed_response['task_breakdown'] if int(task['subtask_id']) == current_subtask_id), None)

                if current_task is None:
                    break 

                print(f"Subtask: {current_task['subtask']}")
                print(f"Agent: {current_task['selected_agent']}")
                print(f"Instruction: {current_task['instruction_to_agent']}")
                print("---")

                agent_output = await self.execute_next_step(current_task)

                print(f"Raw agent output: {agent_output}")
                
                if isinstance(agent_output, dict) and agent_output.get('status') == 'background':
                    self.task_context[str(current_subtask_id)] = {
                        "status": "background",
                        "task_id": agent_output['task_id'],
                        "message": agent_output['message']
                    }
                else:
                    # Mark the previous latest output as not latest
                    for key in self.task_context:
                        if isinstance(self.task_context[key], dict) and self.task_context[key].get('is_latest'):
                            self.task_context[key]['is_latest'] = False
                    
                    self.task_context[str(current_subtask_id)] = agent_output

                # Check for completed background tasks
                completed_tasks = [task_id for task_id, thread in self.background_tasks.items() if not thread.is_alive()]
                for task_id in completed_tasks:
                    print(f"Background task {task_id} completed.")
                    del self.background_tasks[task_id]

                # If the current task was a background task, wait for it to complete
                if isinstance(agent_output, dict) and agent_output.get('status') == 'background':
                    task_id = agent_output['task_id']
                    while task_id in self.background_tasks:
                        await asyncio.sleep(1)
                    print(f"Background task {task_id} completed.")

                continue_task = [
                    "The previous step has been completed. Please analyze the result and determine the next action.",
                    f"Current Subtask ID(THE PROVIDED SUBTASK WITH THIS ID WILL BE EXECUTED): {current_subtask_id}",
                    f"Task context: {json.dumps(self.task_context)}"
                ]
                follow_up_response = await self.generate_response(continue_task)
                parsed_response = self.parse_json_response(follow_up_response)

                if not parsed_response:
                    return "I apologize, but I encountered an error while processing the follow-up. Could you please rephrase your request?"

                continue_execution = parsed_response.get('continue_execution', False)
                
                if parsed_response.get('user_request_fully_finished', False):
                    logging.info("User request fully finished, clearing task context")
                    
                    # Append to memory before clearing
                    if MemoryManager.memory_enabled:
                        memory_content = (
                            "TASK COMPLETION\n" 
                            f"Chat History:\n{self.get_chat_history_as_string()}\n\n"
                            f"Task Context:\n{self.format_task_context()}\n\n"
                            f"Parsed Response:\n{self.format_parsed_response(parsed_response)}"
                        )
                        self.memory_manager.append_to_full_memory("TASK COMPLETION", memory_content)
                        logging.info("Appended task completion to memory")

                    self.task_context.clear()
                    print("Task context cleared as the user request is fully finished.")
                    current_subtask_id = 1
                    print("Resetting current_subtask_id to 1 as task context is cleared.")

                current_subtask_id += 1  # Move to the next subtask

            return parsed_response.get('response_to_user', "I'm sorry, but I couldn't generate a proper follow-up response. Could you please try again?")

        except Exception as e:
            print(f"Unexpected error in process_user_request: {str(e)}")
            return "I'm sorry, but an unexpected error occurred. Please try again or contact support if the issue persists."
        
    def set_update_plan_callback(self, callback):
        self.update_plan_callback = callback

        
    def get_chat_history_as_string(self):
        chat_string = ""
        for i, msg in enumerate(self.chat_history):
            try:
                role = msg.role
                if hasattr(msg.parts[0], 'text'):
                    content = msg.parts[0].text
                else:
                    content = str(msg.parts[0])
                chat_string += f"{role}: {content}\n"
            except Exception as e:
                logging.error(f"Error processing message {i}: {str(e)}")
                logging.error(f"Message type: {type(msg)}")
                logging.error(f"Message content: {str(msg)}")
        return chat_string
    

    def format_task_context(self):
        return "\n".join([f"{key}: {value}" for key, value in self.task_context.items()])

    def format_parsed_response(self, parsed_response):
        return "\n".join([f"{key}: {value}" for key, value in parsed_response.items()])

    async def execute_next_step(self, next_step):
        instruction_with_context = {
            "instruction": next_step['instruction_to_agent'],
            "information_needed": next_step['information_needed'],
            "which_function": next_step['which_function'],
            "task_context": self.task_context
        }
        agent_result = await self.execute_agent_action(next_step['selected_agent'], instruction_with_context)
        
        if isinstance(agent_result, dict) and agent_result.get('status') == 'background':
            task_id = agent_result['thread_id']
            self.background_tasks[task_id] = threading.Thread()
            self.background_tasks[task_id].ident = task_id
            return {"status": "background", "task_id": task_id, "message": f"Task {task_id} started in the background."}
        
        # Limit function output length
        if isinstance(agent_result, dict) and 'result' in agent_result:
            result_str = json.dumps(agent_result['result'])
            if len(result_str) > self.function_output_limit:
                summarized_result = await self.length_limiter.summarize(result_str, self.function_output_limit)
                agent_result['result'] = json.loads(summarized_result)
        
        return {
            "result": agent_result,
            "is_latest": True
        }

    async def execute_agent_action(self, agent, instruction_with_context):
        agent_prompt = self.agent_prompts.get(agent.replace(" ", ""))
        if not agent_prompt:
            return "The selected agent is not defined. Please provide specific instructions or handle this task manually."

        try:
            model = genai.GenerativeModel(
                model_name="models/gemini-1.5-pro",
                generation_config=self.generation_config_1_5,
                safety_settings=self.safety_settings,
                system_instruction=agent_prompt
            )
            agent_chat = model.start_chat(history=[])

            instruction_content = (
                f"INSTRUCTION: {instruction_with_context['instruction']}\n"
                f"INFORMATION NEEDED: {instruction_with_context['information_needed']}\n"
                f"FUNCTION TO USE: {instruction_with_context['which_function']}\n"
                f"Task context: {json.dumps(instruction_with_context['task_context'])}"
            )
            response = await agent_chat.send_message_async(
                content=instruction_content,
                generation_config=self.generation_config_1_5,
                safety_settings=self.safety_settings
            )

            print(f"Raw agent response: {response.text}")

            if not response.text:
                raise ValueError("Empty response from agent.")
            
            parsed_response = self.parse_json_response(response.text)
            if not parsed_response:
                raise ValueError("Failed to parse agent response.")
            
            execution_result = await execute_agent_instruction(json.dumps(parsed_response), agent, self.config)
            if asyncio.iscoroutine(execution_result):
                execution_result = await execution_result
            parsed_response['execution_result'] = execution_result
            
            return parsed_response
        except Exception as e:
            print(f"Error in execute_agent_action for {agent}: {str(e)}")
            return {"error": f"Failed to execute agent action: {str(e)}"}
    
    @property
    def generation_config_1_5(self):
        return genai.types.GenerationConfig(
            temperature=1,
            max_output_tokens=3000,
        )
    
    @property
    def generation_config_1_0(self):
        return genai.types.GenerationConfig(
            temperature=1,
            max_output_tokens=3000,
        )
    
    @property
    def safety_settings(self):
        return {
            HarmCategory.HARM_CATEGORY_HATE_SPEECH: HarmBlockThreshold.BLOCK_NONE,
            HarmCategory.HARM_CATEGORY_HARASSMENT: HarmBlockThreshold.BLOCK_NONE,
            HarmCategory.HARM_CATEGORY_SEXUALLY_EXPLICIT: HarmBlockThreshold.BLOCK_NONE,
            HarmCategory.HARM_CATEGORY_DANGEROUS_CONTENT: HarmBlockThreshold.BLOCK_NONE,
        }

async def main():
    assistant = MainAssistant()
    print("Welcome to the Atlas. Type 'exit' to end the conversation.")
    try:
        while True:
            user_message = input("\nYou: ")
            if user_message.lower() == 'exit':
                print("Thank you for using the Atlas. Goodbye!")
                break
            response = await assistant.process_user_request(user_message)
            print("\nAssistant:", response)
    finally:
        await cleanup_agents()

if __name__ == "__main__":
    asyncio.run(main())